import csv
import math
import os

from openpyxl import Workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_FILE = os.path.join(BASE_DIR, "data", "arduino_data.csv")
DASHBOARD_FILE = os.path.join(
    BASE_DIR,
    "dashboard",
    "Smart_Security_Auto_Dashboard.xlsx",
)

HEADERS = [
    "Date",
    "Time",
    "Access",
    "Temperature_C",
    "Humidity_percent",
    "Distance_cm",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")


def read_csv():
    """Read the saved records. An absent CSV is treated as an empty dataset."""
    if not os.path.exists(CSV_FILE):
        return []

    with open(CSV_FILE, "r", newline="", encoding="utf-8-sig") as file:
        return list(csv.DictReader(file))


def number(value):
    """Convert a valid sensor value to float; keep blanks/errors as missing data."""
    if value in (None, "", "ERROR"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def style_header(worksheet):
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 22


def autosize_columns(worksheet, max_width=30):
    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[letter].width = min(max(longest + 2, 12), max_width)


def add_sensor_chart(dashboard, chart_data, title, value_column, unit, anchor, row_count):
    if row_count == 0:
        return

    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.y_axis.title = unit
    chart.x_axis.title = "Record number"
    chart.legend = None
    chart.height = 7.2
    chart.width = 13.5

    values = Reference(
        chart_data,
        min_col=value_column,
        min_row=1,
        max_row=row_count + 1,
    )
    categories = Reference(
        chart_data,
        min_col=4,
        min_row=2,
        max_row=row_count + 1,
    )
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.x_axis.tickLblSkip = max(1, math.ceil(row_count / 8))
    chart.display_blanks = "gap"
    dashboard.add_chart(chart, anchor)


def build_dashboard():
    rows = read_csv()
    os.makedirs(os.path.dirname(DASHBOARD_FILE), exist_ok=True)

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"
    dashboard = workbook.create_sheet("Dashboard")
    chart_data = workbook.create_sheet("ChartsData")

    # ---------------- Data sheet ----------------
    data_sheet.append(HEADERS)
    for row in rows:
        data_sheet.append([
            row.get("Date", ""),
            row.get("Time", ""),
            row.get("Access", ""),
            number(row.get("Temperature_C")),
            number(row.get("Humidity_percent")),
            number(row.get("Distance_cm")),
        ])

    style_header(data_sheet)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions
    for col, width in {
        "A": 14, "B": 12, "C": 14, "D": 19, "E": 20, "F": 16
    }.items():
        data_sheet.column_dimensions[col].width = width

    for row in data_sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (4, 5, 6) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0"

    # ---------------- Summary metrics ----------------
    granted = sum(1 for row in rows if row.get("Access") == "GRANTED")
    denied = sum(1 for row in rows if row.get("Access") == "DENIED")

    latest_temperature = None
    latest_humidity = None
    latest_distance = None
    for row in reversed(rows):
        if latest_temperature is None:
            latest_temperature = number(row.get("Temperature_C"))
        if latest_humidity is None:
            latest_humidity = number(row.get("Humidity_percent"))
        if latest_distance is None:
            latest_distance = number(row.get("Distance_cm"))
        if all(v is not None for v in (
            latest_temperature, latest_humidity, latest_distance
        )):
            break

    # ---------------- Dashboard layout ----------------
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:Q1")
    dashboard["A1"] = "SMART SECURITY & ENVIRONMENT MONITORING"
    dashboard["A1"].font = Font(size=18, bold=True, color="1F1F1F")
    dashboard["A1"].alignment = Alignment(vertical="center")
    dashboard.row_dimensions[1].height = 32

    metrics = [
        ("A3", "TOTAL ATTEMPTS", len(rows), "B3"),
        ("A4", "GRANTED", granted, "B4"),
        ("A5", "DENIED", denied, "B5"),
        ("D3", "LATEST TEMPERATURE", latest_temperature, "E3"),
        ("D4", "LATEST HUMIDITY", latest_humidity, "E4"),
        ("D5", "LATEST DISTANCE", latest_distance, "E5"),
    ]

    for label_cell, label, value, value_cell in metrics:
        dashboard[label_cell] = label
        dashboard[label_cell].font = Font(bold=True)
        dashboard[label_cell].fill = LABEL_FILL
        dashboard[value_cell] = value if value is not None else "—"
        dashboard[value_cell].font = Font(bold=True, size=12)
        dashboard[value_cell].alignment = Alignment(horizontal="center")
        if isinstance(value, (int, float)) and label.startswith("LATEST"):
            dashboard[value_cell].number_format = "0.0"

    dashboard["F3"] = "°C"
    dashboard["F4"] = "%"
    dashboard["F5"] = "cm"

    for col, width in {
        "A": 22, "B": 14, "C": 4, "D": 25, "E": 14, "F": 8,
        "G": 4, "H": 4, "I": 4, "J": 12, "K": 12, "L": 12, "M": 12,
        "N": 12, "O": 12, "P": 12, "Q": 12,
    }.items():
        dashboard.column_dimensions[col].width = width

    # ---------------- Chart source data ----------------
    chart_data.append(["Access", "Count"])
    chart_data.append(["GRANTED", granted])
    chart_data.append(["DENIED", denied])
    chart_data["D1"] = "Record"
    chart_data["E1"] = "Temperature_C"
    chart_data["F1"] = "Humidity_percent"
    chart_data["G1"] = "Distance_cm"
    style_header(chart_data)

    for record_number, row in enumerate(rows, start=1):
        chart_data.append([
            None,
            None,
            None,
            record_number,
            number(row.get("Temperature_C")),
            number(row.get("Humidity_percent")),
            number(row.get("Distance_cm")),
        ])

    # ---------------- Charts ----------------
    if granted + denied > 0:
        pie = PieChart()
        pie.title = "Access Attempts"
        pie.height = 8
        pie.width = 13.5
        pie_values = Reference(chart_data, min_col=2, min_row=1, max_row=3)
        pie_labels = Reference(chart_data, min_col=1, min_row=2, max_row=3)
        pie.add_data(pie_values, titles_from_data=True)
        pie.set_categories(pie_labels)
        dashboard.add_chart(pie, "A8")

    add_sensor_chart(
        dashboard, chart_data, "Temperature", 5, "°C", "J8", len(rows)
    )
    add_sensor_chart(
        dashboard, chart_data, "Humidity", 6, "%", "A24", len(rows)
    )
    add_sensor_chart(
        dashboard, chart_data, "Distance", 7, "cm", "J24", len(rows)
    )

    chart_data.sheet_state = "hidden"
    workbook.save(DASHBOARD_FILE)
    print(f"Dashboard updated successfully: {DASHBOARD_FILE}")


if __name__ == "__main__":
    build_dashboard()
